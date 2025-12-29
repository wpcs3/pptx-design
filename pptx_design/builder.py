"""
PresentationBuilder - Unified API for creating PowerPoint presentations.

Provides a simple, fluent interface for building presentations using
template master layouts.
"""

import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from pptx import Presentation
from pptx.util import Inches, Pt

from .registry import TemplateRegistry

logger = logging.getLogger(__name__)


class PresentationBuilder:
    """
    Fluent API for building PowerPoint presentations.

    Example:
        builder = PresentationBuilder("consulting_toolkit")
        builder.add_title_slide("Quarterly Review", "Q4 2025")
        builder.add_agenda(["Overview", "Analysis", "Next Steps"])
        builder.add_content_slide("Key Findings", bullets=["Finding 1", "Finding 2"])
        builder.save("output.pptx")
    """

    def __init__(
        self,
        template: str,
        registry: Optional[TemplateRegistry] = None
    ):
        """
        Initialize builder with a template.

        Args:
            template: Template name (e.g., "consulting_toolkit") or path to .pptx file
            registry: Optional TemplateRegistry instance (auto-loaded if not provided)
        """
        self.registry = registry or TemplateRegistry()
        self._template_name = template
        self._template_path = self._resolve_template(template)
        self._slides: List[Dict[str, Any]] = []
        self._metadata: Dict[str, Any] = {
            "company_name": "Company Name",
            "author": "",
            "date": "",
        }

        # Load template info from registry
        self._template_info = self.registry.get_template(template)
        if self._template_info:
            logger.info(f"Loaded template: {template} ({len(self._template_info.get('layouts', []))} layouts)")

    def _resolve_template(self, template: str) -> Path:
        """Resolve template name or path to actual file path."""
        # Check if it's a direct path
        if Path(template).suffix == ".pptx":
            path = Path(template)
            if path.exists():
                return path

        # Try to resolve from registry
        template_info = self.registry.get_template(template)
        if template_info and "path" in template_info:
            path = Path(template_info["path"])
            if path.exists():
                return path

        # Search in templates directory
        templates_dir = Path(__file__).parent.parent / "pptx_templates"
        for pptx_file in templates_dir.rglob("*.pptx"):
            if template.lower() in pptx_file.stem.lower():
                return pptx_file

        raise FileNotFoundError(f"Template not found: {template}")

    # ==================== Metadata Methods ====================

    def set_company(self, name: str) -> "PresentationBuilder":
        """Set company name for footer."""
        self._metadata["company_name"] = name
        return self

    def set_author(self, author: str) -> "PresentationBuilder":
        """Set presentation author."""
        self._metadata["author"] = author
        return self

    def set_date(self, date: str) -> "PresentationBuilder":
        """Set presentation date."""
        self._metadata["date"] = date
        return self

    # ==================== Slide Addition Methods ====================

    def add_title_slide(
        self,
        title: str,
        subtitle: str = "",
        layout: str = "Frontpage"
    ) -> "PresentationBuilder":
        """
        Add a title/cover slide.

        Args:
            title: Main presentation title
            subtitle: Subtitle or tagline
            layout: Layout name (default: "Frontpage")
        """
        self._slides.append({
            "type": "title_slide",
            "layout": layout,
            "content": {
                "title": title,
                "subtitle": subtitle,
            }
        })
        return self

    def add_section_divider(
        self,
        title: str,
        layout: str = "Section breaker"
    ) -> "PresentationBuilder":
        """
        Add a section divider slide.

        Args:
            title: Section title
            layout: Layout name (default: "Section breaker")
        """
        self._slides.append({
            "type": "section_divider",
            "layout": layout,
            "content": {
                "title": title,
            }
        })
        return self

    def add_agenda(
        self,
        items: List[str],
        title: str = "Agenda",
        layout: str = "Agenda"
    ) -> "PresentationBuilder":
        """
        Add an agenda slide.

        Args:
            items: List of agenda items
            title: Slide title (default: "Agenda")
            layout: Layout name (default: "Agenda")
        """
        self._slides.append({
            "type": "agenda",
            "layout": layout,
            "content": {
                "title": title,
                "body": items,
            }
        })
        return self

    def add_content_slide(
        self,
        title: str,
        body: str = "",
        bullets: List[str] = None,
        subtitle: str = "",
        layout: str = "Default"
    ) -> "PresentationBuilder":
        """
        Add a content slide with title and body text or bullets.

        Args:
            title: Slide title
            body: Body text (paragraph)
            bullets: List of bullet points
            subtitle: Optional subtitle
            layout: Layout name (default: "Default")
        """
        content = {"title": title}
        if subtitle:
            content["subtitle"] = subtitle
        if body:
            content["body"] = body
        if bullets:
            content["bullets"] = bullets

        self._slides.append({
            "type": "content",
            "layout": layout,
            "content": content,
        })
        return self

    def add_two_column(
        self,
        title: str,
        left_header: str,
        left_bullets: List[str],
        right_header: str,
        right_bullets: List[str],
        layout: str = "1/2 grey"
    ) -> "PresentationBuilder":
        """
        Add a two-column comparison slide.

        Args:
            title: Slide title
            left_header: Left column header
            left_bullets: Left column bullet points
            right_header: Right column header
            right_bullets: Right column bullet points
            layout: Layout name (default: "1/2 grey")
        """
        self._slides.append({
            "type": "two_column",
            "layout": layout,
            "content": {
                "title": title,
                "left_column": {
                    "header": left_header,
                    "bullets": left_bullets,
                },
                "right_column": {
                    "header": right_header,
                    "bullets": right_bullets,
                }
            }
        })
        return self

    def add_metrics(
        self,
        title: str,
        metrics: List[Dict[str, str]],
        layout: str = "Default"
    ) -> "PresentationBuilder":
        """
        Add a key metrics slide with KPI boxes.

        Args:
            title: Slide title
            metrics: List of {"label": "...", "value": "..."} dicts
            layout: Layout name (default: "Default")

        Example:
            builder.add_metrics("Key Metrics", [
                {"label": "Revenue", "value": "$1.2M"},
                {"label": "Growth", "value": "25%"},
            ])
        """
        self._slides.append({
            "type": "key_metrics",
            "layout": layout,
            "content": {
                "title": title,
                "metrics": metrics,
            }
        })
        return self

    def add_table(
        self,
        title: str,
        headers: List[str],
        data: List[List[str]],
        layout: str = "Default"
    ) -> "PresentationBuilder":
        """
        Add a table slide.

        Args:
            title: Slide title
            headers: Column headers
            data: Table data (list of rows)
            layout: Layout name (default: "Default")

        Example:
            builder.add_table("Comparison",
                headers=["Feature", "Plan A", "Plan B"],
                data=[
                    ["Price", "$100", "$200"],
                    ["Users", "10", "50"],
                ]
            )
        """
        self._slides.append({
            "type": "table_slide",
            "layout": layout,
            "content": {
                "title": title,
                "headers": headers,
                "data": data,
            }
        })
        return self

    def add_chart(
        self,
        title: str,
        chart_type: str,
        categories: List[str],
        series: List[Dict[str, Any]],
        layout: str = "Default"
    ) -> "PresentationBuilder":
        """
        Add a chart slide.

        Args:
            title: Slide title
            chart_type: "column", "bar", "line", "pie"
            categories: X-axis labels
            series: List of {"name": "...", "values": [...]} dicts
            layout: Layout name (default: "Default")

        Example:
            builder.add_chart("Revenue by Quarter",
                chart_type="column",
                categories=["Q1", "Q2", "Q3", "Q4"],
                series=[{"name": "2024", "values": [100, 120, 140, 160]}]
            )
        """
        self._slides.append({
            "type": "data_chart",
            "layout": layout,
            "content": {
                "title": title,
                "chart_data": {
                    "type": chart_type,
                    "categories": categories,
                    "series": series,
                }
            }
        })
        return self

    def add_blank(self, layout: str = "Blank") -> "PresentationBuilder":
        """Add a blank slide."""
        self._slides.append({
            "type": "blank",
            "layout": layout,
            "content": {}
        })
        return self

    def add_end_slide(
        self,
        title: str = "Thank You",
        subtitle: str = "",
        layout: str = "End"
    ) -> "PresentationBuilder":
        """Add an end/closing slide."""
        self._slides.append({
            "type": "end_slide",
            "layout": layout,
            "content": {
                "title": title,
                "subtitle": subtitle,
            }
        })
        return self

    # ==================== Build Methods ====================

    def build(self) -> Presentation:
        """
        Build the presentation and return the Presentation object.

        Returns:
            pptx.Presentation object
        """
        from pptx_extractor.template_generator import TemplateGenerator

        generator = TemplateGenerator(self._template_path, clear_slides=True)

        for slide_spec in self._slides:
            layout_name = slide_spec.get("layout", "Default")
            content = slide_spec.get("content", {})

            # Add company name to footer if not specified
            if "company" not in content and "footer" not in content:
                content["footer"] = self._metadata.get("company_name", "")

            generator.create_slide(layout_name, content)

        return generator.prs

    def save(self, output_path: Union[str, Path]) -> Path:
        """
        Build and save the presentation to a file.

        Args:
            output_path: Path to save the .pptx file

        Returns:
            Path to the saved file
        """
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        prs = self.build()
        prs.save(str(output_path))

        logger.info(f"Saved presentation: {output_path}")
        return output_path

    def preview(self) -> str:
        """
        Get a text preview of the presentation structure.

        Returns:
            Human-readable outline of the presentation
        """
        lines = [f"Presentation: {self._template_name}"]
        lines.append(f"Template: {self._template_path.name}")
        lines.append(f"Slides: {len(self._slides)}")
        lines.append("-" * 40)

        for i, slide in enumerate(self._slides, 1):
            slide_type = slide.get("type", "unknown")
            title = slide.get("content", {}).get("title", "(no title)")
            lines.append(f"  {i}. [{slide_type}] {title}")

        return "\n".join(lines)

    def __len__(self) -> int:
        """Return the number of slides."""
        return len(self._slides)

    def __repr__(self) -> str:
        return f"PresentationBuilder(template='{self._template_name}', slides={len(self._slides)})"


# Convenience function
def create_presentation(template: str) -> PresentationBuilder:
    """
    Create a new presentation builder.

    Args:
        template: Template name or path

    Returns:
        PresentationBuilder instance
    """
    return PresentationBuilder(template)
